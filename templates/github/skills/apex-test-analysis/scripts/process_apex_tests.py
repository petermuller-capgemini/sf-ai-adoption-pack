import csv
import subprocess
import datetime
import json
import os
import glob

# New imports for Excel + plotting
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

# -------- Config --------
INPUT_CSV = 'code-quality/classes.csv'
OUTPUT_DIR = 'test_results'
WAIT_SECS = '10'
EXCEL_PATH = 'code-quality/classes_results.xlsx'
GRAPH_DIR = 'code-quality/graphs'

# Prefer using --json to ensure stdout is JSON.
USE_STDOUT_JSON = True
# ------------------------

today = datetime.date.today().isoformat()

def safe_get_result(obj):
    """
    sf CLI with --json typically wraps the payload under 'result'.
    In other modes, payload keys may be at the top level (as in the sample).
    """
    if not isinstance(obj, dict):
        return {}
    return obj.get('result', obj)

def parse_json_stdout_or_file(stdout_text: str, output_dir: str) -> dict:
    """
    Try parsing JSON from stdout. If that fails, try to locate a JSON file in output_dir.
    Returns a dict (possibly empty) but never raises.
    """
    # 1) Try stdout directly
    if stdout_text:
        try:
            return json.loads(stdout_text)
        except json.JSONDecodeError:
            pass

    # 2) Fallback to latest JSON file in output_dir
    try:
        pattern = os.path.join(output_dir, '*.json')
        candidates = sorted(glob.glob(pattern), key=os.path.getmtime, reverse=True)
        for path in candidates:
            try:
                with open(path, 'r', encoding='utf-8') as f:
                    return json.load(f)
            except Exception:
                continue
    except Exception:
        pass

    return {}

def extract_metrics(payload: dict, class_name: str):
    """
    From test run payload, extract:
      - passing methods count
      - failing methods count
      - coveredPercent for the specific class_name from coverage.coverage array
    """
    res = safe_get_result(payload)

    # Summary fields
    summary = res.get('summary', {}) or {}
    passing = summary.get('passing', 0)
    failing = summary.get('failing', 0)

    # Coverage array and per-class coverage
    coverage_section = res.get('coverage', {}) or {}
    coverage_array = coverage_section.get('coverage', []) or []

    covered_percent = 'N/A'
    target = (class_name or '').strip()

    for item in coverage_array:
        if isinstance(item, dict) and (item.get('name', '').strip() == target):
            cp = item.get('coveredPercent')
            if cp is not None:
                covered_percent = str(cp)  # store as string; will convert later for plotting
            break

    return passing, failing, covered_percent

def ensure_new_columns(fieldnames, new_cols):
    """Ensure we don't duplicate columns if script is re-run the same day."""
    if fieldnames is None:
        fieldnames = []
    for col in new_cols:
        if col not in fieldnames:
            fieldnames.append(col)
    return fieldnames

def save_excel_and_graphs(csv_path: str, excel_path: str, graph_dir: str):
    """
    Read the updated CSV, write a rich Excel file, generate graphs, and embed them.
    """
    os.makedirs(os.path.dirname(excel_path), exist_ok=True)
    os.makedirs(graph_dir, exist_ok=True)

    # Load CSV into DataFrame
    df = pd.read_csv(csv_path)

    # Identify date-stamped coverage columns (YYYY-MM-DD-coverage)
    coverage_cols = [c for c in df.columns if c.endswith('-coverage')]
    if not coverage_cols:
        print("[WARN] No coverage columns found; skipping Excel/graph creation.")
        return

    # Create a wide coverage DF with numeric conversion
    df_wide = df[['ClassName'] + coverage_cols].copy()
    for c in coverage_cols:
        df_wide[c] = pd.to_numeric(df_wide[c], errors='coerce')  # 'N/A' -> NaN

    # Create a long (tidy) DF: ClassName | date | coverage
    long_df = df_wide.melt(
        id_vars=['ClassName'],
        value_vars=coverage_cols,
        var_name='date_col',
        value_name='coverage'
    )
    # Extract the YYYY-MM-DD date out of 'YYYY-MM-DD-coverage'
    long_df['date'] = pd.to_datetime(long_df['date_col'].str.slice(0, 10), errors='coerce')
    long_df = long_df.drop(columns=['date_col']).sort_values(['ClassName', 'date'])

    # Compute overall daily average coverage across classes (ignore NaN)
    # Convert coverage_cols into a sorted date index
    date_index = pd.to_datetime([c[:10] for c in coverage_cols], errors='coerce')
    # Ensure order mirrors the columns order
    overall_series = df_wide[coverage_cols].mean(axis=0, skipna=True)
    overall_series.index = date_index
    overall_series = overall_series.sort_index()

    # ---- Plot 1: Overall coverage over time ----
    overall_png = os.path.join(graph_dir, 'coverage_progress_overall.png')
    plt.figure(figsize=(10, 5))
    plt.plot(overall_series.index, overall_series.values, marker='o', linewidth=2, color='#0078D4')
    plt.title('Overall Code Coverage Progress (Average of Classes)')
    plt.xlabel('Date')
    plt.ylabel('Coverage (%)')
    plt.grid(True, linestyle='--', alpha=0.4)
    plt.tight_layout()
    plt.savefig(overall_png, dpi=150)
    plt.close()

    # # ---- Plot 2: Per-class coverage over time ----
    # # Construct a per-class time series from long_df
    # by_class_png = os.path.join(graph_dir, 'coverage_progress_by_class.png')
    # plt.figure(figsize=(12, 7))
    # for cls, grp in long_df.groupby('ClassName'):
    #     grp = grp.sort_values('date')
    #     plt.plot(grp['date'], grp['coverage'], marker='o', linewidth=1.8, label=cls)
    # plt.title('Per-Class Code Coverage Progress')
    # plt.xlabel('Date')
    # plt.ylabel('Coverage (%)')
    # plt.legend(loc='best', fontsize=8)
    # plt.grid(True, linestyle='--', alpha=0.4)
    # plt.tight_layout()
    # plt.savefig(by_class_png, dpi=150)
    # plt.close()

    # ---- Write Excel with multiple sheets ----
    with pd.ExcelWriter(excel_path, engine='openpyxl', mode='w') as writer:
        df.to_excel(writer, sheet_name='Data', index=False)
        df_wide.to_excel(writer, sheet_name='Coverage (wide)', index=False)
        long_df.to_excel(writer, sheet_name='Coverage (long)', index=False)

    # ---- Embed the images into a 'Charts' sheet ----
    try:
        wb = load_workbook(excel_path)
        if 'Charts' in wb.sheetnames:
            ws = wb['Charts']
            # Clear previous content by recreating the sheet (optional)
            wb.remove(ws)
            ws = wb.create_sheet('Charts')
        else:
            ws = wb.create_sheet('Charts')

        # Place images
        if os.path.exists(overall_png):
            img1 = XLImage(overall_png)
            img1.anchor = 'A1'
            ws.add_image(img1)
        # if os.path.exists(by_class_png):
        #     img2 = XLImage(by_class_png)
        #     # Place second image below the first; adjust row as needed
        #     img2.anchor = 'A30'
        #     ws.add_image(img2)

        wb.save(excel_path)
    except Exception as e:
        print(f"[WARN] Could not embed images into Excel: {e}")

    print(f"[OK] Excel saved: {excel_path}")
    # print(f"[OK] Graphs saved: {overall_png}, {by_class_png}")

def main():
    # 1) Read all rows once (preserve order + existing data)
    with open(INPUT_CSV, mode='r', newline='', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        original_fieldnames = list(reader.fieldnames or [])
        rows = [r for r in reader]

    # 2) Prepare new columns
    passing_col = f'{today}-passing_methods'
    failing_col = f'{today}-failing_methods'
    coverage_col = f'{today}-coverage'
    fieldnames = ensure_new_columns(
        original_fieldnames.copy(),
        [passing_col, failing_col, coverage_col]
    )

    print(f"Total rows to process: {len(rows)}")
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # 3) Process each row
    for idx, row in enumerate(rows):
        class_name = (row.get('ClassName') or '').strip()
        test_class_name = (row.get('TestClassName') or '').strip()

        # Default values in case of issues
        row[passing_col] = row.get(passing_col, '') or '0'
        row[failing_col] = row.get(failing_col, '') or '0'
        row[coverage_col] = row.get(coverage_col, '') or 'N/A'

        if not test_class_name:
            print(f"[WARN] Row {idx+1}: Missing TestClassName for ClassName '{class_name}'. Skipping.")
            continue

        print(f"→ Running tests for: {test_class_name} (target coverage class: {class_name})")

        # Build command
        if USE_STDOUT_JSON:
            test_cmd = [
                'sf', 'apex', 'run', 'test',
                '--class-names', test_class_name,
                '--wait', WAIT_SECS,
                '--output-dir', OUTPUT_DIR,
                '--code-coverage',
                '--json'  # guarantees JSON to stdout
            ]
        else:
            test_cmd = [
                'sf', 'apex', 'run', 'test',
                '--class-names', test_class_name,
                '--result-format', 'json',
                '--wait', WAIT_SECS,
                '--output-dir', OUTPUT_DIR,
                '--code-coverage'
            ]

        # Execute and parse
        try:
            completed = subprocess.run(test_cmd, capture_output=True, text=True, check=True)
            payload = parse_json_stdout_or_file(completed.stdout, OUTPUT_DIR)
            passing, failing, covered_percent = extract_metrics(payload, class_name)
        except subprocess.CalledProcessError as e:
            print(f"[ERROR] Test run failed for {test_class_name}. Return code: {e.returncode}. Error: {e.stdout}")
            payload = parse_json_stdout_or_file(e.stdout or '', OUTPUT_DIR)
            passing, failing, covered_percent = extract_metrics(payload, class_name)
        except Exception as e:
            print(f"[ERROR] Unexpected error for {test_class_name}: {e}")
            passing, failing, covered_percent = 0, 0, 'N/A'

        # Update row with results for today's columns
        row[passing_col] = str(passing)
        row[failing_col] = str(failing)
        row[coverage_col] = str(covered_percent)

    # 4) Write back to CSV preserving order
    with open(INPUT_CSV, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)

    print(f"[OK] Results appended to {INPUT_CSV} with date-stamped columns: {passing_col}, {failing_col}, {coverage_col}")

    # 5) Create Excel and graphs
    save_excel_and_graphs(INPUT_CSV, EXCEL_PATH, GRAPH_DIR)

if __name__ == '__main__':
    main()